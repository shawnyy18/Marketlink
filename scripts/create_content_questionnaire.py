#!/usr/bin/env python3
"""Create the simplified Marketlink website content questionnaire (FSK Codehouse template)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
DOCX_PATH = OUT_DIR / "Marketlink Website Content Questionnaire.docx"
MD_PATH = OUT_DIR / "Marketlink Website Content Questionnaire.md"
LOGO = ROOT / "mklinkname.png"

# Drop an FSK Codehouse logo here to have it used in the header automatically.
FSK_LOGO_CANDIDATES = [
    ROOT / "fsk-logo.png",
    ROOT / "fsk.png",
    ROOT / "public" / "assets" / "fsk-logo.png",
]

COMPANY = "FSK CODEHOUSE"
COMPANY_FULL = "FSK Codehouse"
COMPANY_TAGLINE = "Web design and development"

GREEN = "174B33"
DARK_GREEN = "0C2F21"
MID_GREEN = "357B53"
PALE_GREEN = "EDF4EE"
INK = "17201B"
MUTED = "5F6962"
LINE = "CBD8CE"
WHITE = "FFFFFF"

# FSK Codehouse brand
FSK_INK = "1A1A1A"
FSK_AMBER = "FCA812"
FSK_BLUE = "1F66ED"


SECTIONS = [
    {
        "title": "1. About your business",
        "intro": "These answers shape every headline and description on the site.",
        "draft": None,
        "questions": [
            "What is your exact business name as it should appear on the website?",
            "In one or two sentences, what does Marketlink do?",
            "When did you start, and why?",
            "Who owns or runs the business? Which names and roles can we publish?",
            "Who are your main customers today? List them in order of importance (for example: sari-sari stores, groceries, supermarkets, online sellers, walk-in customers).",
            "Do you sell to businesses only, to individuals, or both?",
            "Which cities or provinces do you actually serve right now?",
            "What is the one thing a visitor should understand about you in five seconds?",
        ],
    },
    {
        "title": "2. Your suppliers",
        "intro": "We need this so the website does not claim more than what is agreed.",
        "draft": "Marketlink Distribution Inc. distributes selected products sourced from established suppliers, including YanYan International.",
        "questions": [
            "What exactly is your relationship with YanYan International?",
            "Which of these words can we use: distributor, reseller, partner, supplier, dealer, exclusive distributor? Only choose what is confirmed in writing.",
            "Are you an exclusive distributor? If yes, for which products, which area, and until when?",
            "Can we name YanYan International on the site, or should we only say 'trusted suppliers'?",
            "Do you carry products from other suppliers? If yes, which ones can we name?",
            "Is there any wording your supplier has asked you to avoid?",
        ],
        "options": ["Keep as written", "Edit wording", "Rewrite", "Remove"],
    },
    {
        "title": "3. Homepage headline",
        "intro": "This is the first thing visitors read.",
        "draft": "PAMPANGA PRODUCT DISTRIBUTION\n\nQuality products. Ready for your shelves.\n\nMarketlink connects stores, resellers, and business buyers with a broad selection of dependable everyday products.\n\nButtons: Browse products / Wholesale inquiries",
        "questions": [
            "Is 'Pampanga product distribution' correct? If not, what should it say?",
            "What promise can you honestly make to customers?",
            "How would you say 'Quality products. Ready for your shelves.' in your own words?",
            "What should the main button do: browse products, call, message, or request a quote?",
            "Which three to five products best represent Marketlink?",
        ],
        "options": ["Keep", "Small edits", "Rewrite", "Remove"],
    },
    {
        "title": "4. Products",
        "intro": "Tell us what buyers need to see before they contact you.",
        "draft": "Current categories: Snacks; Beverages; Instant Beverages; Cookies & Wafers; Chocolates; Bread.",
        "questions": [
            "Are all products shown on the site actually available? List anything to remove or hide.",
            "Which products should appear first?",
            "Are the six categories correct? Give us your preferred names and order.",
            "Should prices be shown? If yes, retail, wholesale, or inquire-only?",
            "Should we show minimum order quantities? If yes, give the details.",
            "Who will update the product list after launch?",
        ],
        "options": ["Approved", "Needs corrections", "Send more info", "Do not publish yet"],
    },
    {
        "title": "5. Wholesale and resellers",
        "intro": "We want the site to describe how you really take orders.",
        "draft": "Current benefits: broad multi-category selection; wholesale and reseller inquiries welcome; product availability confirmed before ordering; one inquiry for multiple product lines.",
        "questions": [
            "Who can order at wholesale prices?",
            "Walk us through your actual process, from first inquiry to confirmed order.",
            "Is there a minimum order? If yes, what is it?",
            "How do buyers get wholesale prices: price list, quotation, or a phone call?",
            "What payment methods and terms can we publish?",
            "Do you offer pickup, delivery, or shipping? To which areas?",
            "Why should a store buy from you instead of somewhere else?",
        ],
        "options": ["Accurate", "Partly accurate", "Rewrite", "Remove section"],
    },
    {
        "title": "6. Your story",
        "intro": "This is the most important section. Rough notes are fine; we will polish the writing.",
        "draft": "Marketlink Distribution Inc. supplies quality consumer products for retail and business buyers. Our catalog includes selected products sourced from established suppliers such as YanYan International.",
        "questions": [
            "How did Marketlink start? What made you begin?",
            "What does the name 'Marketlink' mean to you?",
            "What moments in the business are you proud of?",
            "Where do you want the business to be in three to five years?",
            "Should we mention your warehouse, office, team, or delivery vehicles? Please send photos if yes.",
            "How should the About section sound: formal, warm and family-run, ambitious, or community-focused?",
            "Tell your story in your own words. Write as much or as little as you like.",
        ],
        "options": ["Keep as written", "Use as a base", "Replace with our story", "Remove"],
        "long_response": True,
    },
    {
        "title": "7. Why customers choose you",
        "intro": "Every claim on the site needs something real behind it.",
        "draft": "Current points: Trusted sourcing; Wide selection; Wholesale-friendly; Responsive process.",
        "questions": [
            "What are the top three reasons customers buy from you?",
            "What proof can we show: years in business, number of stores served, warehouse size, permits, or awards?",
            "What service can you promise every single time?",
            "Can we publish customer testimonials? If yes, please send the quotes and names.",
            "What should the website never say about your business?",
        ],
        "options": ["Approved", "Revise claims", "Add proof", "Remove section"],
    },
    {
        "title": "8. Contact details",
        "intro": "Only share details you are happy to make public.",
        "draft": None,
        "questions": [
            "Which email address should receive website inquiries?",
            "Which phone, Viber, or WhatsApp number should we show?",
            "What is your complete office or warehouse address? Should we show a map?",
            "What are your business hours?",
            "Which social media pages are official? Please paste the exact links.",
            "Who answers inquiries, and how fast can you realistically reply?",
            "How would you prefer customers to reach you: website form, call, Messenger, Viber, or email?",
        ],
        "options": ["Keep current form", "Change fields", "Use another channel", "Do not launch form yet"],
    },
    {
        "title": "9. Look and feel",
        "intro": "The site currently uses the green from your logo.",
        "draft": "Current direction: clean, modern, professional, product-focused; forest-green palette based on the supplied Marketlink logo.",
        "questions": [
            "Is the green colour right? Do you have official brand colours?",
            "Are the logo files you sent final?",
            "What real photos can you provide: warehouse, team, delivery vehicles, stocked shelves, products?",
            "Do you have permission to publish every person and store shown in those photos?",
            "Share two or three websites you like, and tell us what you like about them.",
        ],
        "options": ["Approved", "Adjust colours", "Adjust tone", "Send new references"],
    },
    {
        "title": "10. Footer and legal",
        "intro": "Small details, but visitors and partners rely on them.",
        "draft": "Supplier acknowledgment: Marketlink Distribution Inc. distributes selected products sourced from established suppliers, including YanYan International.",
        "questions": [
            "What business name and copyright line should appear at the bottom of every page?",
            "Should the supplier note stay, change, or be removed?",
            "Do you need a Privacy Policy or Terms page? Who will provide the text?",
            "What would a customer type into Google to find you?",
            "What web address (domain) will the site use?",
            "Who gives final approval on everything in this document?",
            "When do you want the website to go live?",
        ],
        "options": ["Approved", "Needs legal review", "Needs supplier review", "Rewrite"],
    },
]


PRIORITY_QUESTIONS = [
    "Which three parts of the website matter most to fix first?",
    "Is anything on the current site wrong or risky?",
    "Which answers do you still need to check with someone else?",
    "Who will send us the final, agreed answers?",
]


def find_fsk_logo() -> Path | None:
    for path in FSK_LOGO_CANDIDATES:
        if path.exists():
            return path
    return None


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=160, start=180, bottom=160, end=180) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int], indent=120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def set_paragraph_border(paragraph, color=LINE, size="6", edge="bottom") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(run) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Heading 1", 17, GREEN, 18, 9),
        ("Heading 2", 13, GREEN, 14, 7),
        ("Heading 3", 11, DARK_GREEN, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_response_lines(doc: Document, count=2) -> None:
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True
    set_font(label.add_run("YOUR ANSWER"), size=7.5, color=MID_GREEN, bold=True)
    for index in range(count):
        line = doc.add_paragraph(" ")
        line.paragraph_format.space_after = Pt(6.5)
        line.paragraph_format.line_spacing = 1.0
        line.paragraph_format.keep_together = True
        line.paragraph_format.keep_with_next = index < count - 1
        set_paragraph_border(line)


def add_current_draft(doc: Document, text: str, label_text="WHAT THE WEBSITE SAYS NOW") -> None:
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(8)
    label.paragraph_format.space_after = Pt(5)
    set_font(label.add_run(label_text), size=8, color=MID_GREEN, bold=True)
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360], indent=120)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE_GREEN)
    set_cell_margins(cell, 180, 220, 180, 220)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, block in enumerate(text.split("\n\n")):
        if index:
            paragraph.add_run("\n\n")
        run = paragraph.add_run(block)
        set_font(run, size=9.5, color=DARK_GREEN, italic=True)


def add_decision_options(doc: Document, options: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    set_font(paragraph.add_run("TICK ONE:  "), size=8, color=DARK_GREEN, bold=True)
    for index, option in enumerate(options):
        if index:
            paragraph.add_run("     ")
        run = paragraph.add_run(f"☐ {option}")
        set_font(run, size=8.5, color=INK)


def add_question(doc: Document, number: int, text: str, long=False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(f"Q{number}. "), size=10, color=GREEN, bold=True)
    set_font(paragraph.add_run(text), size=10, color=INK, bold=True)
    add_response_lines(doc, 4 if long else 2)


def add_cover(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(34)
    paragraph.add_run().add_picture(str(LOGO), width=Inches(2.35))

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    set_font(kicker.add_run("WEBSITE CONTENT QUESTIONNAIRE"), size=9, color=MID_GREEN, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run("Tell us about\nyour business"), size=27, color=DARK_GREEN, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_font(subtitle.add_run("Your answers replace the draft wording on the website with your real story and facts."), size=12, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_width(table, [2160, 7200], indent=120)
    for row_index, (label, value) in enumerate((
        ("Prepared for", "Marketlink Distribution Inc."),
        ("Prepared by", COMPANY_FULL),
        ("Who is answering", ""),
        ("Please return by", ""),
    )):
        left, right = table.rows[row_index].cells
        for cell in (left, right):
            set_cell_margins(cell, 110, 150, 110, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, GREEN)
        set_font(left.paragraphs[0].add_run(label.upper()), size=8, color=WHITE, bold=True)
        set_font(right.paragraphs[0].add_run(value or "____________________________________________"), size=9.5, color=INK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(30)
    note.paragraph_format.space_after = Pt(24)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(note.add_run("The wording on the website today is only a draft. Nothing is final until you confirm it here."), size=9, color=MUTED, italic=True)

    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.paragraph_format.space_after = Pt(0)
    fsk_logo = find_fsk_logo()
    if fsk_logo:
        mark.add_run().add_picture(str(fsk_logo), height=Inches(0.24))
        set_font(mark.add_run("  "), size=10)
    set_font(mark.add_run(COMPANY), size=9.5, color=FSK_INK, bold=True)
    set_font(mark.add_run(f"   {COMPANY_TAGLINE}"), size=9, color=MUTED)


def add_intro(doc: Document) -> None:
    heading = doc.add_heading("How to answer", level=1)
    heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph("Write short, plain answers. Bullet points are fine. You do not need to write polished sentences — we will handle the writing.")
    p.paragraph_format.space_after = Pt(10)

    add_current_draft(
        doc,
        "If you do not know an answer, write 'to be confirmed' and the name of the person who does. Please do not approve a claim just because it sounds good.",
        label_text="ONE IMPORTANT RULE",
    )

    doc.add_heading("Who should help", level=2)
    for item in (
        "Owner or manager — your story, and final approval",
        "Sales or operations — products, orders, and delivery",
        "Whoever handles your suppliers — what we may say about YanYan and others",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(item), size=10, color=INK)

    doc.add_heading("Sending it back", level=2)
    p = doc.add_paragraph(f"Please send one copy with everyone's answers combined, plus any photos, logos, or permits you want on the site. If people disagree, agree internally first, then send the final version to {COMPANY_FULL}.")
    p.paragraph_format.space_after = Pt(0)


def add_final_review(doc: Document) -> None:
    heading = doc.add_heading("11. Last few questions", level=1)
    heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph("This tells us where to start work.")
    p.paragraph_format.space_after = Pt(10)
    for index, question in enumerate(PRIORITY_QUESTIONS, start=1):
        add_question(doc, index, question, long=index <= 2)

    doc.add_heading("Progress checklist", level=2)
    rows = [
        "Homepage headline", "Products", "Wholesale", "Your story",
        "Why choose you", "Contact details", "Supplier wording", "Photos and logos",
    ]
    table = doc.add_table(rows=1, cols=3)
    for index, header in enumerate(("Website section", "Status", "Approved by / date")):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell, 120, 140, 120, 140)
        set_font(cell.paragraphs[0].add_run(header), size=8.5, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    for section_name in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, 150, 140, 150, 140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_font(cells[0].paragraphs[0].add_run(section_name), size=9, color=INK, bold=True)
        set_font(cells[1].paragraphs[0].add_run("☐ Not yet  ☐ Needs changes  ☐ Approved"), size=8, color=INK)
        set_font(cells[2].paragraphs[0].add_run(""), size=9, color=INK)
    set_table_width(table, [2760, 3600, 3000], indent=120)

    doc.add_heading("Sign-off", level=2)
    for label in ("Name", "Signature", "Date"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(8)
        set_font(paragraph.add_run(f"{label}:  "), size=9, color=DARK_GREEN, bold=True)
        set_font(paragraph.add_run("____________________________________________________________"), size=9, color=MUTED)


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    content_width = 9360  # twips, matches 6.5in body width

    # ---- Header: FSK Codehouse branding ----
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

    fsk_logo = find_fsk_logo()
    if fsk_logo:
        p.add_run().add_picture(str(fsk_logo), height=Inches(0.2))
        set_font(p.add_run("  "), size=9)
    set_font(p.add_run(COMPANY), size=9, color=FSK_INK, bold=True)
    p.add_run("\t")
    set_font(p.add_run("Marketlink Distribution Inc.  ·  Website content questionnaire"), size=7.5, color=MUTED)
    set_paragraph_border(p, color=FSK_AMBER, size="8", edge="bottom")

    # ---- Footer: FSK Codehouse branding ----
    footer = section.footer
    f = footer.paragraphs[0]
    f.paragraph_format.space_before = Pt(2)
    f.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    f.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    set_paragraph_border(f, color=LINE, size="6", edge="top")

    set_font(f.add_run(f"{COMPANY_FULL}"), size=7.5, color=FSK_INK, bold=True)
    set_font(f.add_run(f"  ·  {COMPANY_TAGLINE}"), size=7.5, color=MUTED)
    f.add_run("\t")
    page_run = f.add_run()
    set_font(page_run, size=7.5, color=MUTED)
    add_page_field(page_run)
    f.add_run("\t")
    set_font(f.add_run("Confidential — prepared for the client named above"), size=7.5, color=MUTED)

    assert content_width  # keeps the intent of the fixed body width explicit


def build_docx() -> None:
    doc = Document()
    configure_styles(doc)
    configure_header_footer(doc)
    add_cover(doc)
    add_intro(doc)

    for content in SECTIONS:
        heading = doc.add_heading(content["title"], level=1)
        heading.paragraph_format.page_break_before = True
        intro = doc.add_paragraph(content["intro"])
        intro.paragraph_format.space_after = Pt(8)
        if content.get("draft"):
            add_current_draft(doc, content["draft"])
        if content.get("options"):
            add_decision_options(doc, content["options"])
        for question_index, question in enumerate(content["questions"], start=1):
            long = bool(content.get("long_response")) and question_index == len(content["questions"])
            add_question(doc, question_index, question, long=long)

    add_final_review(doc)
    doc.core_properties.title = "Marketlink Website Content Questionnaire"
    doc.core_properties.subject = "Client questionnaire for website content"
    doc.core_properties.author = COMPANY_FULL
    doc.core_properties.company = COMPANY_FULL
    doc.core_properties.keywords = "FSK Codehouse, Marketlink, website, content, questionnaire"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def build_markdown() -> None:
    lines = [
        f"<!-- {COMPANY_FULL} · {COMPANY_TAGLINE} -->",
        "",
        "# Marketlink Website Content Questionnaire",
        "",
        f"**Prepared by:** {COMPANY_FULL}  ",
        "**Prepared for:** Marketlink Distribution Inc.",
        "",
        "---",
        "",
        "## How to answer",
        "",
        "Write short, plain answers. Bullet points are fine. You do not need to write polished sentences — we will handle the writing.",
        "",
        "> **One important rule:** if you do not know an answer, write *to be confirmed* and the name of the person who does. Please do not approve a claim just because it sounds good.",
        "",
        "**Who should help**",
        "",
        "- Owner or manager — your story, and final approval",
        "- Sales or operations — products, orders, and delivery",
        "- Whoever handles your suppliers — what we may say about YanYan and others",
        "",
        f"**Sending it back:** please send one copy with everyone's answers combined, plus any photos, logos, or permits you want on the site. Send the final version to {COMPANY_FULL}.",
        "",
    ]
    for content in SECTIONS:
        lines.extend(["---", "", f"## {content['title']}", "", content["intro"], ""])
        if content.get("draft"):
            lines.extend(["**What the website says now:**", "", f"> {content['draft'].replace(chr(10) * 2, ' / ')}", ""])
        if content.get("options"):
            lines.append("**Tick one:** " + "  ".join(f"☐ {option}" for option in content["options"]))
            lines.append("")
        for index, question in enumerate(content["questions"], start=1):
            lines.extend([f"**Q{index}. {question}**", "", "Your answer:", "", ""])
    lines.extend(["---", "", "## 11. Last few questions", "", "This tells us where to start work.", ""])
    for index, question in enumerate(PRIORITY_QUESTIONS, start=1):
        lines.extend([f"**Q{index}. {question}**", "", "Your answer:", "", ""])
    lines.extend([
        "---", "",
        "## Sign-off", "",
        "**Name:** ___________________________________________", "",
        "**Signature:** ______________________________________", "",
        "**Date:** ___________________________________________", "",
        "---", "",
        f"*{COMPANY_FULL} · {COMPANY_TAGLINE} · Confidential — prepared for the client named above*",
        "",
    ])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
    print(DOCX_PATH)
    print(MD_PATH)
